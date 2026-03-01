import json
import os
import secrets
import string
import openpyxl


# Rastgele 6 haneli eşsiz kod üretme fonksiyonu
def generate_unique_code(existing_codes):
    while True:
        code = ''.join(secrets.choice(string.ascii_letters + string.digits) for _ in range(6))
        if code not in existing_codes:
            return code


def create_davetli_listesi_js(veriler):
    davetliler_json = {}
    for veri in veriler:
        # davetliler_json[veri[COL_TOKEN]] = {"İsim": veri[COL_ISIM], "tip": "i" if veri[COL_OKUL] == "İTÜ" else "o"}
        davetliler_json[veri[COL_TOKEN]] = veri[COL_ISIM]

    with open("davetli-listesi.js", "w", encoding="utf-8") as js_dosyasi:
        js_dosyasi.write(f"davetliler = {json.dumps(davetliler_json, indent=4, ensure_ascii=False)}\n")


def read_and_parse_form_yanitlari_xlsx(dosya_adi):
    # Excel dosyasını yükle
    wb = openpyxl.load_workbook(dosya_adi, data_only=True)
    ws = wb.active  # İlk sayfayı seç

    # Sütun başlıklarını al
    basliklar = [cell.value for cell in ws[1]]

    # Aynı başlığa sahip sütunların indekslerini belirle
    sutun_indexleri = {}
    for i, baslik in enumerate(basliklar):
        if baslik:
            if baslik.startswith(COL_ISIM):
                anahtar = COL_ISIM
            elif baslik.lower() == COL_TELEFON.lower():
                anahtar = COL_TELEFON
            elif baslik.startswith(COL_CINSIYET):
                anahtar = COL_CINSIYET
            elif baslik.endswith(COL_EMAIL):
                anahtar = COL_EMAIL
            elif baslik.startswith(COL_EKLEMEK_ISTEDIKLERINIZ):
                anahtar = COL_EKLEMEK_ISTEDIKLERINIZ
            else:
                anahtar = baslik

            if anahtar not in sutun_indexleri:
                sutun_indexleri[anahtar] = []
            sutun_indexleri[anahtar].append(i)

    # Telefon numarası sütunlarını belirle
    telefon_sutunlari = [COL_TELEFON]

    # Satırları işle
    veriler = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        satir_verisi = {}
        for baslik, index_listesi in sutun_indexleri.items():
            degerler = [row[i] for i in index_listesi if row[i] is not None]

            # Telefon numarası ise string formatına çevir
            if baslik in telefon_sutunlari and degerler:
                satir_verisi[baslik] = str(int(degerler[0]))  # Ondalık kısımdan kurtul
            else:
                satir_verisi[baslik] = degerler[0] if degerler else None

        satir_verisi.pop(COL_EKLEMEK_ISTEDIKLERINIZ, None)
        satir_verisi.pop(COL_CINSIYET, None)

        veriler.append(satir_verisi)
    return veriler


def add_tokens_to_veriler(veriler, existing_tokens):
    for veri in veriler:
        token = veri.get(COL_TOKEN, None)
        if not token:
            token = generate_unique_code(existing_tokens)
            veri[COL_TOKEN] = token
        existing_tokens.add(token)
    return veriler


ITULU_SMS = """Sayın {isim} ,
İTÜ İFTARI YAKLAŞIYOR !

Değerli İTÜ'lü, kaydınız alınmıştır. İlginiz için teşekkür ederiz. Kampüse İTÜ kimliğinizle girebilirsiniz. Davetiyenize linkten ulaşabilirsiniz.

https://dvty.tr/i?t={token}

İftarımızda görüşmek üzere. 
"""
DIGER_SMS = """Sayın {isim} ,
İTÜ İFTARI YAKLAŞIYOR !

Değerli misafirimiz kaydınız alınmıştır. İlginiz için teşekkür ederiz. Kampüse giriş kartınıza linkten ulaşabilirsiniz.

https://dvty.tr/?t={token}

İftarımızda görüşmek üzere. 
"""


def add_msg_to_veriler(veriler):
    for veri in veriler:
        if veri[COL_OKUL] == "İTÜ":
            sms_data = ITULU_SMS.format(isim=veri[COL_ISIM], token=veri[COL_TOKEN])
        else:
            sms_data = DIGER_SMS.format(isim=veri[COL_ISIM], token=veri[COL_TOKEN])
        veri[COL_SMS] = sms_data.replace('\n', '\\n\n')
    return veriler


def create_davetliler_xlsx(davetliler, dosya_adi):
    # Excel dosyasını oluştur
    wb = openpyxl.Workbook()
    ws = wb.active

    # Sütun başlıklarını belirle (dictionary'nin key'leri)
    sutun_basliklari = list(davetliler[0].keys())
    ws.append(sutun_basliklari)

    # Verileri ekle
    for davetli in davetliler:
        ws.append([davetli.get(key, "") for key in sutun_basliklari])

    # Yeni Excel dosyasını kaydet
    wb.save(dosya_adi)


def davetli_listede_var_mi(davetli, davetliler):
    for d in davetliler:
        if d.get(COL_ZAMAN_DAMGASI) == davetli.get(COL_ZAMAN_DAMGASI) and d.get(COL_ISIM) == davetli.get(COL_ISIM):
            return True
    return False


def merge_yanitlar_ve_davetliler_verileri(davetliler, mevcut_davetliler):
    # İlk dosyadaki verileri ikinci dosyadaki verilerle birleştir, token'ları koru
    guncellenmis_davetliler = mevcut_davetliler[:]
    for davetli in davetliler:
        # Eğer bu veri zaten mevcut verilerde varsa, token'ı koru
        mevcut_davetli = next(
            (v for v in mevcut_davetliler if
             v.get(COL_ZAMAN_DAMGASI) == davetli.get(COL_ZAMAN_DAMGASI) and v.get(COL_ISIM) == davetli.get(COL_ISIM)),
            None)
        if mevcut_davetli:
            davetli[COL_TOKEN] = mevcut_davetli[COL_TOKEN]
        if not davetli_listede_var_mi(davetli, guncellenmis_davetliler):
            guncellenmis_davetliler.append(davetli)
    return guncellenmis_davetliler


FILE_FORM_YANITLARI = "İtü iftar 26 (Yanıtlar).xlsx"
FILE_DAVETLILER = "itu_iftari_davetliler.xlsx"

COL_TOKEN = "token"
COL_SMS = "sms"
COL_OKUL = "Okulunuz"
COL_ISIM = "Adınız Soyadınız"
COL_CINSIYET = "Size nasıl hitap etmemizi istersiniz?"
COL_TELEFON = "Telefon Numaranız (5xxxxxxxxx)"
COL_ZAMAN_DAMGASI = "Zaman damgası"
COL_STATUS = "Mesaj gönderildi mi?"
COL_EMAIL = "E-posta adresiniz"
COL_EKLEMEK_ISTEDIKLERINIZ = "Eklemek istedikleriniz"

if __name__ == '__main__':
    # İlk dosyayı oku ve parse et
    veriler = read_and_parse_form_yanitlari_xlsx(FILE_FORM_YANITLARI)
    mevcut_veriler = []
    if os.path.exists(FILE_DAVETLILER):
        mevcut_veriler = read_and_parse_form_yanitlari_xlsx(FILE_DAVETLILER)

    guncellenmis_veriler = merge_yanitlar_ve_davetliler_verileri(veriler, mevcut_veriler)
    mevcut_tokenlar = set(veri[COL_TOKEN] for veri in guncellenmis_veriler if COL_TOKEN in veri)

    # Yeni eklenen kayıtlara token atama işlemini gerçekleştir
    guncellenmis_veriler = add_tokens_to_veriler(guncellenmis_veriler, mevcut_tokenlar)
    guncellenmis_veriler = add_msg_to_veriler(guncellenmis_veriler)

    # Güncellenmiş verilerle .js dosyasını ve ikinci Excel dosyasını oluştur
    create_davetli_listesi_js(guncellenmis_veriler)
    create_davetliler_xlsx(guncellenmis_veriler, FILE_DAVETLILER)

    # Sonuçları yazdır
    for veri in guncellenmis_veriler[:10]:
        print(veri[COL_TOKEN], "\t-\t", veri[COL_TELEFON], "\t-\t", veri[COL_OKUL], "\t-\t", veri[COL_ISIM])
