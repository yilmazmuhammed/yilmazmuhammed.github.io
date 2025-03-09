import openpyxl
import secrets
import string


# Rastgele 6 haneli eşsiz kod üretme fonksiyonu
def generate_unique_code(existing_codes):
    while True:
        code = ''.join(secrets.choice(string.ascii_letters + string.digits) for _ in range(6))
        if code not in existing_codes:
            return code


# Excel dosyasını yükle
dosya_adi = "İTÜ İftarı - 2025 (Yanıtlar).xlsx"  # Dosya adını gerektiği gibi değiştir
wb = openpyxl.load_workbook(dosya_adi)
ws = wb.active  # İlk sayfayı seç

# Sütun başlıklarını al
basliklar = [cell.value for cell in ws[1]]

# Kullanılacak sütunları belirle
ad_soyad_sutunlari = ["Adınız Soyadınız"]
ad_soyad_indexleri = [i for i, sutun in enumerate(basliklar) if sutun in ad_soyad_sutunlari]

# Tüm isimleri topla
tum_isimler = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    for index in ad_soyad_indexleri:
        if row[index]:  # Boş olmayan değerleri al
            tum_isimler.add(row[index])

# Eşsiz kodları oluştur
davetliler = {}
existing_codes = set()
for isim in tum_isimler:
    kod = generate_unique_code(existing_codes)
    existing_codes.add(kod)
    davetliler[kod] = isim

# JavaScript dosyasını oluştur
js_icerik = "davetliler = {\n"
js_icerik += ",\n".join(f'  "{kod}": "{isim}"' for kod, isim in davetliler.items())
js_icerik += "\n}\n"

with open("davetli-listesi.js", "w", encoding="utf-8") as js_dosyasi:
    js_dosyasi.write(js_icerik)

print("davetli-listesi.js dosyası oluşturuldu.")
